"""Read-only Ozon Seller API adapter."""
import csv
import io
import re
import urllib.parse
import urllib.request
from datetime import date, datetime, time as dt_time, timedelta, timezone
from openpyxl import load_workbook
from .legacy import Ozon as BaseOzon, ApiError, ALLOWED

ALLOWED.update({
    '/v1/analytics/stocks', '/v3/posting/fbo/list', '/v4/posting/fbs/list',
    '/v1/finance/accrual/by-day', '/v1/finance/accrual/types',
    '/v1/finance/accrual/postings', '/v1/finance/products/buyout',
    '/v1/report/placement/by-products/create', '/v1/report/info',
    '/v1/returns/list', '/v2/product/info/stocks-by-warehouse/fbs',
})

MOSCOW_TZ = timezone(timedelta(hours=3))
POSTING_NUMBER = re.compile(r'^\d+-\d+-\d+$')


class Ozon(BaseOzon):
    def analytics(self, products, stocks, progress):
        skus = sorted({str(p['sku']) for p in products if p.get('sku')} |
                      {str(s['sku']) for s in stocks if s.get('sku')})
        unique = {}
        for offset in range(0, len(skus), 100):
            data = self.call('/v1/analytics/stocks', {'skus': skus[offset:offset + 100]})
            for row in data['items']:
                key = (str(row['sku']), str(row['warehouse_id']))
                if key in unique:
                    raise ApiError('Повтор строки аналитики Ozon. Итоги не рассчитаны.')
                unique[key] = row
            progress(f'Аналитика: {min(offset + 100, len(skus))} / {len(skus)} SKU')
        if products and not unique:
            raise ApiError('Ozon не вернул строки аналитики для каталога.')
        return list(unique.values())

    def fbs_stocks(self, products, progress):
        """Get seller-warehouse (FBS/rFBS) free stock for the current catalog."""
        skus = sorted({str(p.get('sku')) for p in products if p.get('sku')})
        rows = {}
        for offset in range(0, len(skus), 1000):
            chunk = skus[offset:offset + 1000]
            cursor = ''
            seen = set()
            for _ in range(10000):
                data = self.call('/v2/product/info/stocks-by-warehouse/fbs', {
                    'cursor': cursor, 'limit': 1000, 'sku': chunk,
                })
                batch = data.get('products') or []
                if not isinstance(batch, list):
                    raise ApiError('Ozon вернул некорректные FBS-остатки.')
                for row in batch:
                    key = (str(row.get('sku') or ''), str(row.get('warehouse_id') or ''))
                    rows[key] = {
                        'sku': str(row.get('sku') or ''),
                        'offer_id': str(row.get('offer_id') or ''),
                        'warehouse_id': str(row.get('warehouse_id') or ''),
                        'warehouse_name': str(row.get('warehouse_name') or ''),
                        'free_stock': int(row.get('free_stock') or 0),
                        'present': int(row.get('present') or 0),
                        'reserved': int(row.get('reserved') or 0),
                    }
                if not data.get('has_next', False):
                    break
                nxt = data.get('cursor') or ''
                if not nxt or nxt in seen:
                    raise ApiError('Ozon повторил курсор FBS-остатков.')
                seen.add(nxt)
                cursor = nxt
            else:
                raise ApiError('Слишком много страниц FBS-остатков.')
            progress(f'FBS-остатки: {min(offset + 1000, len(skus))} / {len(skus)} SKU')
        result = list(rows.values())
        return {
            'quantity': sum(max(0, int(r.get('free_stock') or 0)) for r in result),
            'rows': result,
        }

    def postings(self, days, progress):
        until = datetime.now(timezone.utc)
        since = until - timedelta(days=days)
        cursor, seen, rows = '', set(), {}
        for _ in range(10000):
            data = self.call('/v3/posting/fbo/list', {
                'cursor': cursor, 'limit': 100, 'sort_dir': 'ASC',
                'filter': {'since': since.isoformat(), 'to': until.isoformat()},
                'with': {'analytics_data': True, 'financial_data': False, 'legal_info': False},
            })
            for posting in data['postings']:
                rows[posting['posting_number']] = {
                    'number': posting['posting_number'], 'status': posting['status'],
                    'created_at': posting['created_at'],
                    'warehouse_id': str((posting.get('analytics_data') or {}).get('warehouse_id') or '') or None,
                    'warehouse_name': (posting.get('analytics_data') or {}).get('warehouse_name') or None,
                    'products': [{'sku': str(p['sku']), 'offer_id': p.get('offer_id', ''),
                                  'name': p.get('name', ''), 'quantity': int(p['quantity'])}
                                 for p in posting['products']],
                }
            progress(f'История FBO: {len(rows)} отправлений')
            if not data.get('has_next', False):
                return {'since': since.isoformat(), 'to': until.isoformat(),
                        'warehouse_attribution': True, 'postings': list(rows.values())}
            nxt = data.get('cursor', '')
            if not nxt or nxt in seen:
                raise ApiError('Повтор курсора истории. Предыдущие данные сохранены.')
            seen.add(nxt)
            cursor = nxt
        raise ApiError('История слишком велика. Уменьшите период.')

    @staticmethod
    def _as_date(value):
        if isinstance(value, date):
            return value
        try:
            return date.fromisoformat(str(value))
        except ValueError as exc:
            raise ApiError('Некорректная дата периода оцифровки.') from exc

    @staticmethod
    def _price(value):
        if isinstance(value, dict):
            value = value.get('amount', value.get('price'))
        try:
            return float(value or 0)
        except (TypeError, ValueError):
            return 0.0

    @staticmethod
    def _report_number(value):
        if isinstance(value, (int, float)) and not isinstance(value, bool):
            return float(value)
        text = str(value or '').strip().replace('\xa0', '').replace(' ', '').replace(',', '.')
        for suffix in ('₽', 'руб.', 'руб'):
            text = text.replace(suffix, '')
        try:
            return float(text)
        except ValueError:
            return 0.0

    @staticmethod
    def _closed_last_day(requested_last):
        today_moscow = datetime.now(MOSCOW_TZ).date()
        if requested_last >= today_moscow:
            return today_moscow - timedelta(days=1), True
        return requested_last, False

    def _orders_for_period(self, first, last, progress):
        if last < first:
            return {'amount': 0.0, 'quantity': 0, 'postings': 0, 'fbo': 0, 'fbs': 0, 'errors': []}
        since_msk = datetime.combine(first, dt_time.min, tzinfo=MOSCOW_TZ)
        until_msk = datetime.combine(last, dt_time.max, tzinfo=MOSCOW_TZ)
        since = since_msk.astimezone(timezone.utc).isoformat()
        until = until_msk.astimezone(timezone.utc).isoformat()
        result = {'amount': 0.0, 'quantity': 0, 'postings': 0, 'fbo': 0, 'fbs': 0, 'errors': []}
        for scheme, path in (('fbo', '/v3/posting/fbo/list'), ('fbs', '/v4/posting/fbs/list')):
            cursor = ''
            seen = set()
            try:
                for _ in range(10000):
                    data = self.call(path, {
                        'cursor': cursor, 'limit': 100, 'sort_dir': 'ASC',
                        'filter': {'since': since, 'to': until},
                        'with': {'analytics_data': False, 'financial_data': False, 'legal_info': False},
                    })
                    postings = data.get('postings') or []
                    if not isinstance(postings, list):
                        raise ApiError(f'Ozon вернул некорректный список заказов {scheme.upper()}.')
                    for posting in postings:
                        result['postings'] += 1
                        result[scheme] += 1
                        for product in posting.get('products') or []:
                            quantity = int(product.get('quantity') or 0)
                            result['quantity'] += quantity
                            result['amount'] += self._price(product.get('price')) * quantity
                    progress(f'Заказы {scheme.upper()}: {result[scheme]} отправлений')
                    if not data.get('has_next', False):
                        break
                    nxt = data.get('cursor') or ''
                    if not nxt or nxt in seen:
                        raise ApiError(f'Ozon повторил курсор заказов {scheme.upper()}.')
                    seen.add(nxt)
                    cursor = nxt
                else:
                    raise ApiError(f'Слишком много страниц заказов {scheme.upper()}.')
            except ApiError as exc:
                result['errors'].append(f'{scheme.upper()}: {exc}')
                progress(f'Заказы {scheme.upper()} недоступны: {exc}')
        result['amount'] = round(result['amount'], 2)
        return result

    def _returns_for_period(self, first, last, progress):
        """Return decisions (returns/refusals) for TrueStats buyout percentage."""
        if last < first:
            return {'quantity': 0, 'rows': 0, 'by_type': {}}
        since_msk = datetime.combine(first, dt_time.min, tzinfo=MOSCOW_TZ)
        until_msk = datetime.combine(last, dt_time.max, tzinfo=MOSCOW_TZ)
        since = since_msk.astimezone(timezone.utc).isoformat()
        until = until_msk.astimezone(timezone.utc).isoformat()
        last_id = 0
        seen = set()
        quantity = 0
        row_count = 0
        by_type = {}
        for _ in range(10000):
            data = self.call('/v1/returns/list', {
                'filter': {
                    'logistic_return_date': {'time_from': since, 'time_to': until},
                },
                'limit': 500,
                'last_id': last_id,
            })
            batch = data.get('returns') or []
            if not isinstance(batch, list):
                raise ApiError('Ozon вернул некорректный список возвратов.')
            for item in batch:
                q = int((item.get('product') or {}).get('quantity') or 0)
                quantity += max(0, q)
                row_count += 1
                kind = str(item.get('type') or 'Unknown')
                by_type[kind] = by_type.get(kind, 0) + max(0, q)
            progress(f'Возвраты/отказы: {row_count} строк · {quantity} шт')
            if not data.get('has_next', False):
                break
            if not batch:
                raise ApiError('Ozon указал следующую страницу возвратов без данных.')
            nxt = batch[-1].get('id')
            try:
                nxt = int(nxt)
            except (TypeError, ValueError):
                raise ApiError('Ozon не вернул last_id для возвратов.') from None
            if nxt in seen:
                raise ApiError('Ozon повторил last_id возвратов.')
            seen.add(nxt)
            last_id = nxt
        else:
            raise ApiError('Слишком много страниц возвратов.')
        return {'quantity': quantity, 'rows': row_count, 'by_type': by_type}

    def _download_report_bytes(self, url):
        url = str(url or '').strip()
        parsed = urllib.parse.urlparse(url)
        host = (parsed.hostname or '').lower()
        if parsed.scheme != 'https' or not (host.endswith('.ozon.ru') or host.endswith('.ozone.ru')):
            raise ApiError('Ozon вернул недопустимую ссылку на отчёт размещения.')
        request = urllib.request.Request(url, headers={'User-Agent': 'ISK/0.1'})
        try:
            with urllib.request.urlopen(request, timeout=40) as response:
                data = response.read(50 * 1024 * 1024 + 1)
        except OSError as exc:
            raise ApiError('Не удалось скачать отчёт о стоимости размещения.') from exc
        if len(data) > 50 * 1024 * 1024:
            raise ApiError('Отчёт о стоимости размещения слишком большой.')
        return data

    def _placement_sum(self, raw):
        target = 'Начисленная стоимость размещения'
        if raw[:2] == b'PK':
            try:
                workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
                for sheet in workbook.worksheets:
                    rows = sheet.iter_rows(values_only=True)
                    index = None
                    for _ in range(30):
                        row = next(rows, None)
                        if row is None:
                            break
                        normalized = [' '.join(str(v or '').split()) for v in row]
                        if target in normalized:
                            index = normalized.index(target)
                            break
                    if index is None:
                        continue
                    total = 0.0
                    for row in rows:
                        if index < len(row):
                            total += self._report_number(row[index])
                    return round(total, 2)
            except Exception as exc:
                raise ApiError('Не удалось прочитать XLSX-отчёт стоимости размещения.') from exc
        else:
            text = None
            for encoding in ('utf-8-sig', 'cp1251'):
                try:
                    text = raw.decode(encoding)
                    break
                except UnicodeDecodeError:
                    continue
            if text is None:
                raise ApiError('Неизвестная кодировка отчёта стоимости размещения.')
            sample = text[:4096]
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=';,\t')
            except csv.Error:
                dialect = csv.excel
                dialect.delimiter = ';'
            rows = csv.reader(io.StringIO(text), dialect)
            index = None
            for _ in range(30):
                row = next(rows, None)
                if row is None:
                    break
                normalized = [' '.join(str(v or '').split()) for v in row]
                if target in normalized:
                    index = normalized.index(target)
                    break
            if index is not None:
                total = 0.0
                for row in rows:
                    if index < len(row):
                        total += self._report_number(row[index])
                return round(total, 2)
        raise ApiError('В отчёте Ozon нет столбца «Начисленная стоимость размещения».')

    def _placement_cost_report(self, first, last, progress):
        if last < first:
            return 0.0
        created = self.call('/v1/report/placement/by-products/create', {
            'date_from': first.isoformat(), 'date_to': last.isoformat(),
        })
        code = created.get('code') or ''
        if not code:
            raise ApiError('Ozon не вернул идентификатор отчёта стоимости размещения.')
        progress('Формируется отчёт стоимости размещения…')
        for _ in range(45):
            info = self.call('/v1/report/info', {'code': code}).get('result') or {}
            status = info.get('status')
            if status == 'success':
                value = self._placement_sum(self._download_report_bytes(info.get('file')))
                progress(f'Стоимость размещения: {value:.2f} ₽')
                return value
            if status == 'failed':
                raise ApiError('Ozon не смог сформировать отчёт стоимости размещения.')
            self.sleep(1)
        raise ApiError('Ozon слишком долго формирует отчёт стоимости размещения.')

    def _posting_accrual_rows(self, posting_numbers, first, last, progress):
        rows = []
        for offset in range(0, len(posting_numbers), 200):
            chunk = posting_numbers[offset:offset + 200]
            data = self.call('/v1/finance/accrual/postings', {'posting_numbers': chunk})
            groups = data.get('posting_accruals') or []
            if not isinstance(groups, list):
                raise ApiError('Ozon вернул некорректные начисления по отправлениям.')
            for group in groups:
                posting_number = str(group.get('posting_number') or '')
                for accrual in group.get('accruals') or []:
                    accrual_date = str(accrual.get('accrual_date') or '')[:10]
                    try:
                        row_date = date.fromisoformat(accrual_date)
                    except ValueError:
                        continue
                    if first <= row_date <= last:
                        rows.append({
                            'posting_number': posting_number,
                            'accrual_date': accrual_date,
                            'type_id': accrual.get('type_id'),
                            'amount': self._price(accrual.get('accrued')),
                            'quantity': self._price(accrual.get('quantity')),
                            'seller_price': self._price(accrual.get('seller_price')),
                            'sku': accrual.get('sku'),
                        })
            progress(f'Детализация начислений: {min(offset + 200, len(posting_numbers))} / {len(posting_numbers)} отправлений')
        return rows

    def _flatten_fee_rows(self, accruals, include_posting):
        result = []
        for accrual in accruals:
            unit = str(accrual.get('unit_number') or '').strip()
            if not include_posting and POSTING_NUMBER.match(unit):
                continue
            accrual_date = str(accrual.get('date') or '')[:10]

            def append_fee(fee, sku=None):
                if not isinstance(fee, dict) or fee.get('type_id') is None:
                    return
                result.append({
                    'posting_number': unit, 'accrual_date': accrual_date,
                    'type_id': fee.get('type_id'), 'amount': self._price(fee.get('accrued')),
                    'quantity': 0, 'seller_price': 0, 'sku': sku,
                })

            append_fee(accrual.get('non_item_fee'))
            for fee in (accrual.get('container_fees') or {}).get('fees') or []:
                append_fee(fee)
            for item in (accrual.get('item_fees') or {}).get('fees') or []:
                sku = item.get('sku')
                for fee in item.get('fees') or []:
                    append_fee(fee, sku)
            for product in (accrual.get('posting') or {}).get('products') or []:
                sku = product.get('sku')
                for fee in (product.get('delivery') or {}).get('services') or []:
                    append_fee(fee, sku)
        return result

    def _sales_facts(self, accruals):
        totals = {'revenue': 0.0, 'partner_programs': 0.0, 'discount_points': 0.0}
        by_sku = {}
        for accrual in accruals:
            for product in (accrual.get('posting') or {}).get('products') or []:
                commission = product.get('commission') or {}
                revenue = self._price(commission.get('sale_amount'))
                partner = self._price(commission.get('coinvestment'))
                bonus = self._price(commission.get('bonus'))
                totals['revenue'] += revenue
                totals['partner_programs'] += partner
                totals['discount_points'] += bonus
                sku = product.get('sku')
                if sku not in (None, '', 0, '0'):
                    key = str(sku)
                    row = by_sku.setdefault(key, {
                        'sku': key, 'revenue': 0.0,
                        'partner_programs': 0.0, 'discount_points': 0.0,
                    })
                    row['revenue'] += revenue
                    row['partner_programs'] += partner
                    row['discount_points'] += bonus
        totals['products'] = list(by_sku.values())
        return totals

    def finance_transactions(self, date_from, date_to, progress):
        first = self._as_date(date_from)
        requested_last = self._as_date(date_to)
        if requested_last < first:
            raise ApiError('Дата окончания не может быть раньше даты начала.')
        if (requested_last - first).days > 30:
            raise ApiError('Финансовый период в одном запросе должен быть не больше 31 дня.')
        effective_last, current_day_excluded = self._closed_last_day(requested_last)

        type_data = self.call('/v1/finance/accrual/types', {})
        accrual_types = type_data.get('accrual_types') or []
        if not isinstance(accrual_types, list):
            raise ApiError('Ozon вернул некорректный справочник начислений.')

        raw_accruals = []
        posting_numbers = set()
        if effective_last >= first:
            current_day = first
            while current_day <= effective_last:
                last_id = ''
                seen = set()
                for _ in range(1000):
                    data = self.call('/v1/finance/accrual/by-day', {
                        'date': current_day.isoformat(), 'last_id': last_id,
                    })
                    batch = data.get('accruals') or []
                    if not isinstance(batch, list):
                        raise ApiError('Ozon вернул некорректный список начислений.')
                    raw_accruals.extend(batch)
                    for item in batch:
                        unit = str(item.get('unit_number') or '').strip()
                        if POSTING_NUMBER.match(unit):
                            posting_numbers.add(unit)
                    nxt = data.get('last_id') or ''
                    if not nxt or not batch:
                        break
                    if nxt in seen:
                        raise ApiError('Ozon повторил курсор финансовых начислений.')
                    seen.add(nxt)
                    last_id = nxt
                else:
                    raise ApiError('Слишком много страниц финансовых начислений за один день.')
                progress(f'Начисления Ozon: {current_day.isoformat()} · {len(raw_accruals)} блоков')
                current_day += timedelta(days=1)

        posting_rows = (
            self._posting_accrual_rows(sorted(posting_numbers), first, effective_last, progress)
            if posting_numbers and effective_last >= first else []
        )
        accrual_rows = list(posting_rows)
        accrual_rows.extend(self._flatten_fee_rows(raw_accruals, include_posting=False))
        by_day_fee_rows = self._flatten_fee_rows(raw_accruals, include_posting=True)
        sales_facts = self._sales_facts(raw_accruals)
        progress(f'Строк начислений для сверки: {len(accrual_rows)}')

        buyout_products = []
        buyout_error = None
        if effective_last >= first:
            try:
                data = self.call('/v1/finance/products/buyout', {
                    'date_from': first.isoformat(), 'date_to': effective_last.isoformat(),
                })
                buyout_products = data.get('products') or []
                if not isinstance(buyout_products, list):
                    raise ApiError('Ozon вернул некорректный отчёт по выкупленным товарам.')
                progress(f'Выкупы маркетплейсом: {len(buyout_products)} строк')
            except ApiError as exc:
                buyout_error = str(exc)
                progress(f'Выкупы маркетплейсом недоступны: {exc}')

        placement_cost = None
        placement_error = None
        if effective_last >= first:
            try:
                placement_cost = self._placement_cost_report(first, effective_last, progress)
            except ApiError as exc:
                placement_error = str(exc)
                progress(f'Стоимость размещения недоступна: {exc}')
        else:
            placement_cost = 0.0

        returns = None
        returns_error = None
        if effective_last >= first:
            try:
                returns = self._returns_for_period(first, effective_last, progress)
            except ApiError as exc:
                returns_error = str(exc)
                progress(f'Возвраты/отказы недоступны: {exc}')

        orders = self._orders_for_period(first, effective_last, progress)

        return {
            'since': first.isoformat(), 'to': requested_last.isoformat(),
            'effective_to': effective_last.isoformat() if effective_last >= first else None,
            'current_day_excluded': current_day_excluded,
            'accrual_types': accrual_types, 'accrual_rows': accrual_rows,
            'by_day_fee_rows': by_day_fee_rows, 'sales_facts': sales_facts,
            'buyout_products': buyout_products, 'buyout_error': buyout_error,
            'placement_cost': placement_cost, 'placement_error': placement_error,
            'returns': returns, 'returns_error': returns_error,
            'orders': orders,
        }
